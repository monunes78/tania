"""
Parsers de documentos: PDF, DOCX, XLSX, TXT.
Retorna texto limpo extraído do arquivo.
"""
import io
import structlog

log = structlog.get_logger()


def parse_pdf(data: bytes) -> str:
    import pdfplumber
    texts = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for i, page in enumerate(pdf.pages):
            text = page.extract_text()
            if text:
                texts.append(text.strip())
    return "\n\n".join(texts)


def parse_docx(data: bytes) -> str:
    from docx import Document
    doc = Document(io.BytesIO(data))
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    return "\n\n".join(paragraphs)


def parse_xlsx(data: bytes) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    rows = []
    for sheet in wb.worksheets:
        rows.append(f"=== Planilha: {sheet.title} ===")
        for row in sheet.iter_rows(values_only=True):
            row_text = "\t".join(str(c) if c is not None else "" for c in row)
            if row_text.strip():
                rows.append(row_text)
    return "\n".join(rows)


def parse_txt(data: bytes) -> str:
    for enc in ("utf-8", "latin-1", "cp1252"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


PARSERS = {
    "pdf": parse_pdf,
    "docx": parse_docx,
    "xlsx": parse_xlsx,
    "txt": parse_txt,
}


def extract_text(data: bytes, file_type: str) -> str:
    parser = PARSERS.get(file_type.lower())
    if not parser:
        raise ValueError(f"Tipo de arquivo não suportado: {file_type}")
    text = parser(data)
    log.info("parser.extracted", file_type=file_type, chars=len(text))
    return text
